"""One-time, resumable import of the XPACE experimental-class Excel history.

Usage (from repository root):
  python scripts/import-xpace-experimental.py "C:\\path\\Graficos de Acompanhamento 2026.xlsx"

The script reads credentials from .env.local, never prints them, and only imports
the tbAgendamento table. It intentionally does not match old names to students or
current classes: the workbook lacks reliable identifiers and class times.
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
from datetime import date, datetime
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import openpyxl


COMPANY_SLUG = "xpace"
SHEET_NAME = "Comparec. Experimental"
TABLE_NAME = "tbAgendamento"


def load_env(root: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    for line in (root / ".env.local").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


class Rest:
    def __init__(self, url: str, key: str):
        self.base = f"{url.rstrip('/')}/rest/v1"
        self.headers = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json", "Accept": "application/json"}

    def call(self, resource: str, method: str = "GET", query: dict[str, str] | None = None, body: object | None = None, single: bool = False):
        target = f"{self.base}/{resource}"
        if query:
            target += "?" + urlencode(query)
        headers = dict(self.headers)
        if method in {"POST", "PATCH"}:
            headers["Prefer"] = "return=representation"
        request = Request(target, data=json.dumps(body).encode("utf-8") if body is not None else None, headers=headers, method=method)
        try:
            with urlopen(request, timeout=40) as response:
                raw = response.read().decode("utf-8")
                payload = json.loads(raw) if raw else None
                return payload[0] if single and isinstance(payload, list) and payload else payload
        except HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"Supabase retornou HTTP {error.code}: {detail[:300]}") from error


def text(value) -> str:
    return "" if value is None else str(value).strip().replace("\xa0", " ")


def iso(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError("A coluna Data contém um valor inválido")


def mapped(value, values: dict[str, str], fallback: str) -> str:
    return values.get(text(value).casefold(), fallback)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python scripts/import-xpace-experimental.py <arquivo.xlsx>")
    source = Path(sys.argv[1]).expanduser().resolve()
    if not source.is_file():
        raise SystemExit("Arquivo Excel não encontrado.")
    root = Path(__file__).resolve().parents[1]
    env = load_env(root)
    url = env.get("NEXT_PUBLIC_SUPABASE_URL", "")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY", "")
    if not url or not key:
        raise SystemExit("NEXT_PUBLIC_SUPABASE_URL ou SUPABASE_SERVICE_ROLE_KEY não configurada em .env.local.")
    api = Rest(url, key)
    # Validate the workbook before creating a resumable batch. This leaves no
    # empty import record if Excel has the source locked.
    book = openpyxl.load_workbook(source, data_only=True, read_only=False)
    sheet = book[SHEET_NAME]
    table = sheet.tables[TABLE_NAME]
    min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(table.ref)
    headers = [sheet.cell(min_row, col).value for col in range(min_col, max_col + 1)]
    companies = api.call("companies", query={"select": "id", "slug": f"eq.{COMPANY_SLUG}", "active": "eq.true"})
    if not companies:
        raise SystemExit("Empresa XPACE não encontrada.")
    company_id = companies[0]["id"]
    # Excel can keep an exclusive binary lock while still allowing openpyxl to read
    # the workbook. Metadata is enough to make this one-time import resumable.
    metadata = source.stat()
    checksum = hashlib.sha256(f"{source.name}:{metadata.st_size}:{metadata.st_mtime_ns}".encode("utf-8")).hexdigest()
    batches = api.call("xpace_lead_import_batches", query={"select": "id", "tenant_company_id": f"eq.{company_id}", "source_checksum": f"eq.{checksum}"})
    if batches:
        batch_id = batches[0]["id"]
        print("Retomando lote de importação existente.")
    else:
        batch = api.call("xpace_lead_import_batches", method="POST", body={"tenant_company_id": company_id, "source_file_name": source.name, "source_sheet_name": SHEET_NAME, "source_checksum": checksum, "note": "Importação do histórico de aulas experimentais do Excel."}, single=True)
        batch_id = batch["id"]
        print("Lote de importação criado.")

    imported = skipped = 0
    for excel_row, values in enumerate(sheet.iter_rows(min_row=min_row + 1, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True), start=min_row + 1):
        row = dict(zip(headers, values))
        if not any(value not in (None, "") for value in values):
            continue
        existing = api.call("xpace_leads", query={"select": "id", "legacy_import_batch_id": f"eq.{batch_id}", "legacy_row_number": f"eq.{excel_row}"})
        if existing:
            skipped += 1
            continue
        attendance = mapped(row.get("Status"), {"compareceu": "COMPARECEU", "faltou": "FALTOU"}, "NAO_INFORMADO")
        enrollment = mapped(row.get("Matriculou?"), {"sim": "MATRICULOU", "não": "NAO_MATRICULOU", "nao": "NAO_MATRICULOU"}, "NAO_INFORMADO")
        stage = "GANHO" if enrollment == "MATRICULOU" else "PERDIDO" if enrollment == "NAO_MATRICULOU" else "AULA_EXPERIMENTAL"
        raw_payload = {str(key): iso(value) if isinstance(value, (datetime, date)) else value for key, value in row.items()}
        lead = api.call("xpace_leads", method="POST", body={"tenant_company_id": company_id, "full_name": text(row.get("Aluno")) or f"LEAD HISTÓRICO {excel_row}", "pipeline_stage": stage, "loss_note": "HISTÓRICO IMPORTADO: NÃO MATRICULOU." if stage == "PERDIDO" else None, "legacy_import_batch_id": batch_id, "legacy_row_number": excel_row, "legacy_payload": raw_payload}, single=True)
        appointment = api.call("xpace_lead_appointments", method="POST", body={"tenant_company_id": company_id, "lead_id": lead["id"], "scheduled_on": iso(row.get("Data")), "booking_kind": mapped(row.get("Tipo de Agendamento"), {"novo": "NOVO", "reagendado": "REAGENDAMENTO", "recuperado": "RECUPERACAO"}, "NOVO"), "confirmation_status": mapped(row.get("Confirmação"), {"sim": "CONFIRMADO", "não": "NAO_CONFIRMADO", "nao": "NAO_CONFIRMADO"}, "NAO_INFORMADO"), "attendance_status": attendance, "enrollment_outcome": enrollment, "attendant_name_snapshot": text(row.get("Atendente")) or None, "modality_name_snapshot": text(row.get("Modalidade")) or None, "instructor_name_snapshot": text(row.get("Professor")) or None, "legacy_week_label": text(row.get("Semana")) or None, "note": text(row.get("Obs")) or None, "survey_status": mapped(row.get("Pesquisa"), {"sim": "ENVIADA", "não": "NAO_ENVIADA", "nao": "NAO_ENVIADA"}, "NAO_INFORMADO")}, single=True)
        api.call("xpace_lead_activities", method="POST", body={"tenant_company_id": company_id, "lead_id": lead["id"], "appointment_id": appointment["id"], "activity_type": "IMPORTADO", "body": f"HISTÓRICO IMPORTADO DA LINHA {excel_row} DA PLANILHA."})
        imported += 1
        if imported % 25 == 0:
            print(f"{imported} linhas importadas nesta execução...")
    api.call("xpace_lead_import_batches", method="PATCH", query={"id": f"eq.{batch_id}"}, body={"imported_rows": imported + skipped})
    print(f"Concluído. Importadas agora: {imported}; já existentes: {skipped}; total do lote: {imported + skipped}.")


if __name__ == "__main__":
    main()
